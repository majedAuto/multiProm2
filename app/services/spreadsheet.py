from __future__ import annotations

import json
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from app.schemas.api import FileType, OutputConfig, RowConfig


def detect_file_type(path: Path, requested: FileType = FileType.auto) -> FileType:
    if requested != FileType.auto:
        return requested
    ext = path.suffix.lower()
    if ext == ".csv":
        return FileType.csv
    if ext in {".xlsx", ".xlsm"}:
        return FileType.xlsx
    raise ValueError(f"Unsupported file extension: {ext}")


def load_dataframe(path: Path, file_type: FileType, sheet_name: Optional[str] = None) -> pd.DataFrame:
    if file_type == FileType.csv:
        return pd.read_csv(path)
    if file_type == FileType.xlsx:
        return pd.read_excel(path, sheet_name=sheet_name or 0, engine="openpyxl")
    raise ValueError(f"Unsupported file type: {file_type}")


def save_dataframe(df: pd.DataFrame, path: Path, file_type: FileType, sheet_name: Optional[str] = None) -> None:
    if file_type == FileType.csv:
        df.to_csv(path, index=False)
        return
    if file_type == FileType.xlsx:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name or "Sheet1")
        return
    raise ValueError(f"Unsupported file type: {file_type}")


def save_processed_output(
    *,
    df: pd.DataFrame,
    source_path: Path,
    output_path: Path,
    file_type: FileType,
    source_sheet_name: Optional[str],
    output_cfg: OutputConfig,
    row_cfg: RowConfig,
    original_columns: list[str],
) -> None:
    if file_type != FileType.xlsx or not _should_use_excel_overlay(output_cfg):
        save_dataframe(df, output_path, file_type=file_type, sheet_name=source_sheet_name)
        return

    workbook = load_workbook(source_path)
    target_sheet_name = output_cfg.target_sheet_name or source_sheet_name or workbook.sheetnames[0]
    ws = workbook[target_sheet_name] if target_sheet_name in workbook.sheetnames else workbook.create_sheet(target_sheet_name)

    export_columns = _resolve_export_columns(df, output_cfg, row_cfg, original_columns)
    if not export_columns:
        # Fallback: write full sheet if no export columns could be derived.
        save_dataframe(df, output_path, file_type=file_type, sheet_name=source_sheet_name)
        return

    start_col = _parse_excel_column(output_cfg.target_start_column) or 1
    header_row = output_cfg.target_start_row or 1
    if header_row < 1:
        header_row = 1

    write_headers = bool(output_cfg.write_headers)
    data_start_row = header_row + 1 if write_headers else header_row

    if write_headers:
        for offset, col_name in enumerate(export_columns):
            ws.cell(row=header_row, column=start_col + offset, value=col_name)

    for row_offset in range(len(df)):
        excel_row = data_start_row + row_offset
        for col_offset, col_name in enumerate(export_columns):
            value = df.at[row_offset, col_name] if col_name in df.columns else None
            ws.cell(row=excel_row, column=start_col + col_offset, value=_excel_safe_value(value))

    workbook.save(output_path)


def _should_use_excel_overlay(output_cfg: OutputConfig) -> bool:
    return any(
        [
            bool(output_cfg.target_sheet_name),
            output_cfg.target_start_row is not None,
            bool(output_cfg.target_start_column),
            bool(output_cfg.export_columns),
        ]
    )


def _parse_excel_column(column_ref: Optional[str]) -> Optional[int]:
    if column_ref is None:
        return None
    ref = str(column_ref).strip()
    if not ref:
        return None
    if ref.isdigit():
        col = int(ref)
        return col if col >= 1 else None
    try:
        return int(column_index_from_string(ref.upper()))
    except Exception:
        return None


def _resolve_export_columns(
    df: pd.DataFrame,
    output_cfg: OutputConfig,
    row_cfg: RowConfig,
    original_columns: list[str],
) -> list[str]:
    if output_cfg.export_columns:
        return list(dict.fromkeys([c for c in output_cfg.export_columns if c]))

    derived: list[str] = []
    for col in [row_cfg.status_column, row_cfg.error_column]:
        if col and col in df.columns:
            derived.append(col)
    if output_cfg.write_raw_output and row_cfg.raw_output_column and row_cfg.raw_output_column in df.columns:
        derived.append(row_cfg.raw_output_column)

    original_set = set(original_columns)
    for col in df.columns:
        if col not in original_set:
            derived.append(col)

    return list(dict.fromkeys(derived))


def _excel_safe_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if pd.isna(value):
        return None
    return value
